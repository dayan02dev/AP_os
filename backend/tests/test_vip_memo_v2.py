"""VIP memo v2 ("Navigator") — read service, exporters and the three surfaces'
GET + download routes. Hermetic: FakeSupabase + a synthetic fixture company
("Acme Robotics"); no real applicant content.
"""
from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.deps import get_current_user
from app.main import app
from app.services import vip_memo, vip_memo_v2, vip_memo_v2_export
from tests.fixtures.fake_supabase import FakeSupabase

FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "vip_memo_v2_fake.json").read_text())
PILOT = sorted(vip_memo.PILOT_APPLICATION_IDS)[0]
OTHER_PILOT = sorted(vip_memo.PILOT_APPLICATION_IDS)[1]
NON_PILOT = "99999999-9999-9999-9999-999999999999"   # a VIP app with a generated memo
UNKNOWN = "11111111-1111-1111-1111-111111111111"     # exists nowhere
TIR_APP = "22222222-2222-2222-2222-222222222222"     # native + effective TIR
MOVED_APP = "33333333-3333-3333-3333-333333333333"   # TIR row moved to VIP
MOVED_OUT = "44444444-4444-4444-4444-444444444444"   # VIP row moved to TIR

SECTION_TITLES = (
    "Deal snapshot",
    "What does this company do",
    "Why this solution matters",
    "The product",
    "Technology edge",
    "Competitive landscape",
    "Addressable market",
    "Founding team",
    "Milestones",
    "Use of funds",
    "Key risks & mitigants",
    "Questions for the founders",
    "IC Reviewer Notes",
)
REVIEWER_AREAS = (
    "Moats & Defensibility", "Red Flags", "Commercial Viability", "Team Assessment",
    "Funding & Structure", "Technical Validation", "Go-to-Market Clarity",
    "IP & Entity Structure", "Follow-on Fundability", "Overall Verdict",
)


def _memo_row(app_id: str, memo: dict | None = None, **over) -> dict:
    return {
        "id": f"memo-{app_id[:4]}", "application_id": app_id, "version": 1,
        "status": "ready", "is_current": True,
        "memo_json": copy.deepcopy(memo if memo is not None else FIXTURE), **over,
    }


def _apps(memo_rows: list[dict]) -> dict:
    return {
        "vip_memos": memo_rows,
        "sip_applications": [{"id": i} for i in (PILOT, OTHER_PILOT, NON_PILOT)]
                            + [{"id": MOVED_OUT, "moved_to_track": "tir"}],
        "tir_applications": [{"id": TIR_APP}, {"id": MOVED_APP, "moved_to_track": "sip"}],
        "vip_memo_exports": [],
    }


@pytest.fixture(autouse=True)
def no_bundled_memos(monkeypatch, tmp_path):
    """Point the bundled-file fallback at an empty dir so a developer's local
    (gitignored) memo files never leak into these tests."""
    empty = tmp_path / "bundled_empty"
    empty.mkdir()
    monkeypatch.setattr(vip_memo_v2, "BUNDLED_DIR", empty)
    return empty


@pytest.fixture
def fake(monkeypatch):
    sb = FakeSupabase(_apps([_memo_row(PILOT), _memo_row(NON_PILOT)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    return sb


# ─── service ────────────────────────────────────────────────────────────

def test_get_memo_v2_returns_stored_memo_for_pilot(fake):
    memo = vip_memo_v2.get_memo_v2(PILOT)
    assert memo is not None
    assert memo["name"] == "Acme Robotics"
    assert memo["version"] == 2


def test_db_memo_served_for_non_pilot(fake):
    memo = vip_memo_v2.get_memo_v2(NON_PILOT)
    assert memo is not None and memo["name"] == "Acme Robotics"


def test_get_memo_v2_missing_row_returns_none(fake):
    assert vip_memo_v2.get_memo_v2(OTHER_PILOT) is None


@pytest.mark.parametrize("over", [{"status": "failed"}, {"status": "running"}, {"is_current": False}])
def test_get_memo_v2_ignores_non_current_or_non_ready_rows(monkeypatch, over):
    sb = FakeSupabase(_apps([_memo_row(PILOT, **over)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    assert vip_memo_v2.get_memo_v2(PILOT) is None


@pytest.mark.parametrize("mutate", [
    lambda m: m.update(version=1),
    lambda m: m.pop("sections"),
    lambda m: m.update(sections="nope"),
    lambda m: m["sections"].pop("risks"),
    lambda m: m.pop("name"),
])
def test_get_memo_v2_rejects_bad_shape(monkeypatch, mutate):
    memo = copy.deepcopy(FIXTURE)
    mutate(memo)
    sb = FakeSupabase(_apps([_memo_row(PILOT, memo)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    assert vip_memo_v2.get_memo_v2(PILOT) is None


def test_get_memo_v2_never_raises(monkeypatch):
    def boom():
        raise RuntimeError("db down")
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", boom)
    assert vip_memo_v2.get_memo_v2(PILOT) is None


def _bundle(directory: Path, app_id: str, memo: dict) -> None:
    (directory / f"{app_id}.json").write_text(json.dumps(memo))


def test_get_memo_v2_falls_back_to_bundled_file(monkeypatch, no_bundled_memos):
    sb = FakeSupabase(_apps([]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    _bundle(no_bundled_memos, PILOT, {**copy.deepcopy(FIXTURE), "name": "Bundled Co"})
    assert vip_memo_v2.get_memo_v2(PILOT)["name"] == "Bundled Co"


def test_get_memo_v2_db_memo_wins_over_bundled_file(fake, no_bundled_memos):
    _bundle(no_bundled_memos, PILOT, {**copy.deepcopy(FIXTURE), "name": "Bundled Co"})
    assert vip_memo_v2.get_memo_v2(PILOT)["name"] == "Acme Robotics"


def test_get_memo_v2_bundled_file_used_when_db_down(monkeypatch, no_bundled_memos):
    def boom():
        raise RuntimeError("db down")
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", boom)
    _bundle(no_bundled_memos, PILOT, copy.deepcopy(FIXTURE))
    assert vip_memo_v2.get_memo_v2(PILOT)["name"] == "Acme Robotics"


def test_bundled_pilot_still_served(monkeypatch, no_bundled_memos):
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: FakeSupabase(_apps([])))
    _bundle(no_bundled_memos, PILOT, {**copy.deepcopy(FIXTURE), "name": "Bundled Co"})
    memo = vip_memo_v2.get_memo_v2(PILOT)
    assert memo["name"] == "Bundled Co"


def test_get_memo_v2_bundled_file_is_validated(fake, no_bundled_memos):
    bad = copy.deepcopy(FIXTURE)
    bad["version"] = 1
    _bundle(no_bundled_memos, OTHER_PILOT, bad)
    (no_bundled_memos / "not-json.json").write_text("{")
    assert vip_memo_v2.get_memo_v2(OTHER_PILOT) is None


def test_get_memo_v2_bundled_file_corrupt_never_raises(monkeypatch, no_bundled_memos):
    sb = FakeSupabase(_apps([]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    (no_bundled_memos / f"{PILOT}.json").write_text("{not json")
    assert vip_memo_v2.get_memo_v2(PILOT) is None


def test_download_filename_is_sanitised():
    assert vip_memo_v2.download_filename({"name": "Acme Robotics"}, "pdf") == "Acme_Robotics_IC_Memo.pdf"
    assert vip_memo_v2.download_filename({"name": 'A/b "c"'}, "docx") == "A_b_c_IC_Memo.docx"
    assert vip_memo_v2.download_filename({}, "pdf") == "VIP_IC_Memo.pdf"


# ─── exporters ──────────────────────────────────────────────────────────

def _docx_text(body: bytes) -> tuple[str, str]:
    doc = Document(io.BytesIO(body))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    header = " ".join(p.text for s in doc.sections for p in s.header.paragraphs)
    return "\n".join(parts), header


def test_render_docx_v2_has_header_and_all_sections():
    body = vip_memo_v2_export.render_docx_v2(FIXTURE)
    assert body[:2] == b"PK"
    text, header = _docx_text(body)
    assert "CONFIDENTIAL" in header
    for title in SECTION_TITLES + REVIEWER_AREAS:
        assert title in text, title
    assert "Acme Robotics" in text
    assert "Acme Robotics advantage" in text
    assert "Reviewer Name" in text and "Signature" in text
    assert "Prepared by: Test fixture team" in text
    # The IC recommendation is left out of the export for now.
    assert "IC recommendation" not in text
    assert "CONDITIONAL APPROVAL" not in text
    assert FIXTURE["recommendation"]["summary"] not in text
    doc = Document(io.BytesIO(body))
    sec = doc.sections[0]
    assert round(sec.page_width.inches, 1) == 8.5
    assert round(sec.left_margin.inches, 1) == 0.8


def test_render_pdf_v2_has_header_and_all_sections():
    body = vip_memo_v2_export.render_pdf_v2(FIXTURE)
    assert body.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(body))
    pages = [p.extract_text() or "" for p in reader.pages]
    assert len(pages) >= 2
    assert all("CONFIDENTIAL" in p for p in pages)
    text = "\n".join(pages)
    for title in SECTION_TITLES + REVIEWER_AREAS:
        assert title in text, title
    assert "Page 1" in pages[0]
    assert "Prepared by: Test fixture team" in text
    # US Letter
    box = reader.pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (612, 792)


# Long text stored as bullet lists: a bullet is a string or [lead, *sub-points];
# risks may be {title, risk, handled} objects.
POINTS = copy.deepcopy(FIXTURE)
POINTS["plain"] = ["Plain point one.", "Plain point two."]
POINTS["sections"]["what"] = {"paragraphs": ["What point one.", "What point two."],
                              "analogy": ["Like a crane.", "But smaller."]}
POINTS["sections"]["why"]["rows"][3] = ["Hard to copy", [["Lead moat.", "Sub moat a."], "Second moat."]]
POINTS["sections"]["product"]["offers"] = [["Fake Arm", ["Gets one arm.", "Gets an app."], ["₹5 L", "Pilot free"]]]
POINTS["sections"]["product"]["note"] = ["Product note one.", "Product note two."]
POINTS["sections"]["market"]["tam_note"] = ["TAM note one.", "TAM note two."]
POINTS["sections"]["team"]["notes"] = ["Team note one.", "Team note two."]
POINTS["sections"]["milestones"]["infra"] = ["Lab space", "GPU time"]
POINTS["sections"]["risks"]["rows"] = [
    {"title": "Regulatory", "risk": ["No licence yet.", "Class unclear."], "handled": ["Test licence held."]},
    ["Legacy risk", "Legacy mitigant"],
]


def test_render_docx_v2_renders_bullet_lists():
    doc = Document(io.BytesIO(vip_memo_v2_export.render_docx_v2(POINTS)))
    bullets = [p.text for p in doc.paragraphs if p.style.name.startswith("List Bullet")]
    for item in ("Plain point two.", "What point two.", "But smaller.", "Lead moat.", "Sub moat a.",
                 "Second moat.", "Product note two.", "TAM note two.", "Team note two.", "GPU time"):
        assert item in bullets, item
    sub = next(p for p in doc.paragraphs if p.text == "Sub moat a.")
    assert sub.style.name == "List Bullet 2"
    text, _ = _docx_text(vip_memo_v2_export.render_docx_v2(POINTS))
    assert "• Gets one arm.\n• Gets an app." in text  # list inside a table cell
    assert "Regulatory" in text and "• No licence yet.\n• Class unclear." in text
    assert "Legacy mitigant" in text
    assert "['" not in text  # no Python list reprs leak into the document


def test_render_pdf_v2_renders_bullet_lists():
    reader = PdfReader(io.BytesIO(vip_memo_v2_export.render_pdf_v2(POINTS)))
    text = "\n".join(p.extract_text() or "" for p in reader.pages)
    for item in ("Plain point two.", "Sub moat a.", "Gets an app.", "TAM note two.", "Class unclear.", "Legacy mitigant"):
        assert item in text, item
    assert "['" not in text


def test_renderers_tolerate_sparse_memo():
    sparse = {"version": 2, "name": "Acme Robotics", "sections": {}}
    assert vip_memo_v2_export.render_docx_v2(sparse)[:2] == b"PK"
    assert vip_memo_v2_export.render_pdf_v2(sparse).startswith(b"%PDF")


# ─── routes ─────────────────────────────────────────────────────────────

@pytest.fixture
def client():
    yield TestClient(app)
    app.dependency_overrides.clear()


def _as(roles, user_id="u-1"):
    app.dependency_overrides[get_current_user] = lambda: {
        "user_id": user_id, "email": "t@example.com", "roles": roles, "track": None,
    }


ROUTES = [
    ("leadership", ["leadership"], "/leadership/applications/{id}/vip-memo-v2"),
    ("admin", ["admin"], "/admin/platform/applications/sip/{id}/vip-memo-v2"),
    ("reviewer", ["reviewer"], "/reviewer/applications/sip/{id}/vip-memo-v2"),
]


@pytest.fixture
def assigned(monkeypatch):
    """Reviewer assignment check: every app is assigned to u-1, nothing to anyone else."""
    from app.services import reviewer_query
    calls = []

    native = {MOVED_APP: "tir", TIR_APP: "tir"}  # everything else is natively "sip"

    def _fetch(uid, track, app_id):
        calls.append((uid, track, app_id))
        return {"application": {}} if uid == "u-1" and track == native.get(app_id, "sip") else None
    monkeypatch.setattr(reviewer_query, "fetch_application_for_reviewer", _fetch)
    return calls


@pytest.mark.parametrize("surface,roles,path", ROUTES)
def test_get_route_returns_memo(client, fake, assigned, surface, roles, path):
    _as(roles)
    r = client.get(path.format(id=PILOT))
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "Acme Robotics"


@pytest.mark.parametrize("surface,roles,path", ROUTES)
def test_get_route_404_when_not_available(client, fake, assigned, surface, roles, path):
    _as(roles)
    for app_id in (UNKNOWN, TIR_APP, OTHER_PILOT):  # unknown, TIR, VIP without a memo
        r = client.get(path.format(id=app_id))
        assert r.status_code == 404
        assert r.json()["detail"]["code"] == "memo_not_available"


@pytest.mark.parametrize("surface,roles,path", ROUTES)
@pytest.mark.parametrize("fmt,magic,ext", [("pdf", b"%PDF", "pdf"), ("docx", b"PK", "docx")])
def test_download_route(client, fake, assigned, surface, roles, path, fmt, magic, ext):
    _as(roles)
    r = client.post(path.format(id=PILOT) + f"/download?format={fmt}")
    assert r.status_code == 200, r.text
    assert r.content.startswith(magic)
    assert f'filename="Acme_Robotics_IC_Memo.{ext}"' in r.headers["content-disposition"]


@pytest.mark.parametrize("surface,roles,path", ROUTES)
def test_download_route_404_and_bad_format(client, fake, assigned, surface, roles, path):
    _as(roles)
    assert client.post(path.format(id=OTHER_PILOT) + "/download?format=pdf").status_code == 404
    assert client.post(path.format(id=PILOT) + "/download?format=exe").status_code == 422


def test_reviewer_route_enforces_assignment(client, fake, assigned):
    _as(["reviewer"], user_id="someone-else")
    r = client.get(f"/reviewer/applications/sip/{PILOT}/vip-memo-v2")
    assert r.status_code == 404
    r = client.post(f"/reviewer/applications/sip/{PILOT}/vip-memo-v2/download?format=pdf")
    assert r.status_code == 404
    assert ("someone-else", "sip", PILOT) in assigned


def test_admin_route_rejects_tir_track(client, fake):
    _as(["admin"])
    r = client.get(f"/admin/platform/applications/tir/{TIR_APP}/vip-memo-v2")
    assert r.status_code == 404


@pytest.mark.parametrize("roles,path", [
    (["applicant"], f"/leadership/applications/{PILOT}/vip-memo-v2"),
    (["applicant"], f"/admin/platform/applications/sip/{PILOT}/vip-memo-v2"),
    (["applicant"], f"/reviewer/applications/sip/{PILOT}/vip-memo-v2"),
])
def test_routes_require_capability(client, fake, roles, path):
    _as(roles)
    assert client.get(path).status_code == 403


# ─── generated-memo serving (any VIP app) ────────────────────────────────

def test_v2_get_404_for_unknown_or_tir_app(client, monkeypatch, assigned):
    sb = FakeSupabase(_apps([_memo_row(TIR_APP), _memo_row(MOVED_OUT)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    for _surface, roles, path in ROUTES:
        _as(roles)
        for app_id in (UNKNOWN, TIR_APP, MOVED_OUT):  # MOVED_OUT: effective track is TIR
            r = client.get(path.format(id=app_id))
            assert r.status_code == 404, (path, app_id)
            assert r.json()["detail"]["code"] == "memo_not_available"


def test_moved_tir_app_is_served_on_every_surface(client, monkeypatch, assigned):
    sb = FakeSupabase(_apps([_memo_row(MOVED_APP)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    for _surface, roles, path in ROUTES:
        _as(roles)
        assert client.get(path.format(id=MOVED_APP)).status_code == 200, path
        assert client.post(path.format(id=MOVED_APP) + "/download?format=pdf").status_code == 200, path
    # the reviewer's assignment is looked up under the NATIVE track
    assert ("u-1", "tir", MOVED_APP) in assigned


def test_reviewer_moved_app_any_url_track_and_assignment_still_enforced(client, monkeypatch, assigned):
    sb = FakeSupabase(_apps([_memo_row(MOVED_APP)]))
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    _as(["reviewer"])
    for url_track in ("sip", "tir"):
        assert client.get(f"/reviewer/applications/{url_track}/{MOVED_APP}/vip-memo-v2").status_code == 200
    _as(["reviewer"], user_id="someone-else")
    assert client.get(f"/reviewer/applications/sip/{MOVED_APP}/vip-memo-v2").status_code == 404


@pytest.mark.parametrize("surface,roles,path", ROUTES)
def test_recommendation_not_in_any_get(client, fake, assigned, surface, roles, path):
    assert "recommendation" in FIXTURE  # the stored memo does carry one
    _as(roles)
    r = client.get(path.format(id=NON_PILOT))
    assert r.status_code == 200
    assert "recommendation" not in r.json()
    assert "CONDITIONAL APPROVAL" not in r.text
    assert FIXTURE["recommendation"]["summary"] not in r.text
    assert "recommendation" in vip_memo_v2.get_memo_v2(NON_PILOT)  # service keeps it; public_view strips


def test_public_view_does_not_mutate():
    memo = {"name": "x", "recommendation": {"verdict": "APPROVE"}}
    assert vip_memo_v2.public_view(memo) == {"name": "x"}
    assert "recommendation" in memo


def _with_storage(sb, files: dict):
    class Bucket:
        def __init__(self, name):
            self.name = name

        def download(self, path):
            return files[(self.name, path)]

    class Storage:
        def from_(self, name):
            return Bucket(name)

    sb.storage = Storage()


def test_stored_export_preferred(client, monkeypatch, assigned):
    rows = _apps([_memo_row(NON_PILOT)])
    rows["vip_memo_exports"] = [
        {"memo_id": f"memo-{NON_PILOT[:4]}", "format": "pdf", "bucket": "vip-memos", "path": "a/b.pdf"},
    ]
    sb = FakeSupabase(rows)
    _with_storage(sb, {("vip-memos", "a/b.pdf"): b"%PDF-stored"})
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    monkeypatch.setattr(vip_memo_v2_export, "render", lambda *a, **k: pytest.fail("rendered despite stored export"))
    mid = f"memo-{NON_PILOT[:4]}"
    assert vip_memo_v2.stored_export(mid, "pdf") == b"%PDF-stored"
    assert vip_memo_v2.stored_export(mid, "docx") is None  # no docx export row
    for _surface, roles, path in ROUTES:
        _as(roles)
        r = client.post(path.format(id=NON_PILOT) + "/download?format=pdf")
        assert r.status_code == 200 and r.content == b"%PDF-stored", path
        assert r.headers["content-type"] == "application/pdf"


def test_stored_export_missing_or_broken_falls_back_to_render(fake, monkeypatch):
    mid = f"memo-{NON_PILOT[:4]}"
    assert vip_memo_v2.stored_export(mid, "pdf") is None
    body, media = vip_memo_v2.export_for(mid, FIXTURE, "pdf")
    assert body.startswith(b"%PDF") and media == "application/pdf"
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: (_ for _ in ()).throw(RuntimeError("down")))
    assert vip_memo_v2.stored_export(mid, "pdf") is None


def test_bundled_memo_never_gets_a_db_export(client, monkeypatch, no_bundled_memos):
    """DB has only a non-current row with an export; the served memo is the bundled
    file, so that export must NOT be served."""
    rows = _apps([_memo_row(PILOT, is_current=False)])
    rows["vip_memo_exports"] = [{"memo_id": f"memo-{PILOT[:4]}", "format": "pdf",
                                 "bucket": "vip-memos", "path": "old.pdf"}]
    sb = FakeSupabase(rows)
    _with_storage(sb, {("vip-memos", "old.pdf"): b"%PDF-stale"})
    monkeypatch.setattr(vip_memo_v2, "get_admin_client", lambda: sb)
    _bundle(no_bundled_memos, PILOT, {**copy.deepcopy(FIXTURE), "name": "Bundled Co"})
    memo, memo_id = vip_memo_v2.load_memo_v2(PILOT)
    assert memo["name"] == "Bundled Co" and memo_id is None
    _as(["leadership"])
    r = client.post(f"/leadership/applications/{PILOT}/vip-memo-v2/download?format=pdf")
    assert r.status_code == 200 and r.content != b"%PDF-stale" and r.content.startswith(b"%PDF")


def test_pdf_replacement_logs_count_only(caplog):
    with caplog.at_level("WARNING"):
        assert vip_memo_v2_export._pdf_safe("a 日本 b") == "a ?? b"
    assert "2 character" in caplog.text and "日本" not in caplog.text


def test_v1_still_pilot_gated(client, fake):
    for roles, path in ((["leadership"], f"/leadership/applications/{NON_PILOT}/vip-memo"),
                        (["admin"], f"/admin/platform/applications/sip/{NON_PILOT}/vip-memo")):
        _as(roles)
        r = client.post(path)
        assert r.status_code == 404, (path, r.text)
        assert r.json()["detail"]["code"] == "vip_memo_pilot_only"


# ─── exporter fixes ──────────────────────────────────────────────────────

def _both_texts(memo: dict) -> list[str]:
    pdf = PdfReader(io.BytesIO(vip_memo_v2_export.render_pdf_v2(memo)))
    pdf_text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    docx_text, _ = _docx_text(vip_memo_v2_export.render_docx_v2(memo))
    return [pdf_text, docx_text]


def test_export_unmatched_why_label_rendered():
    memo = copy.deepcopy(FIXTURE)
    memo["sections"]["why"]["rows"].append(["Moonshot angle", "Fake unlisted reason."])
    for text in _both_texts(memo):
        assert "Moonshot angle" in text and "Fake unlisted reason." in text
        assert "Status quo" in text  # matched labels still grouped under their titles


def test_export_competitor_stage_and_funds_columns():
    for text in _both_texts(FIXTURE):
        assert "Series A" in text  # 5th competitor column
        assert "₹50 L" in text or "Rs 50 L" in text  # funds amount by column name
        assert "Fake salaries" in text  # funds key items
    blocks = vip_memo_v2_export.build_blocks(FIXTURE)
    funds = next(b for b in blocks if b[0] == "table" and b[1][0] == "Category")
    assert funds[2][0][:3] == ["Team", funds[2][0][1], "50%"]
    assert funds[2][-1][2] == "100%"  # computed 50+30+20


def test_export_funds_total_pct_only_when_rows_parse():
    memo = copy.deepcopy(FIXTURE)
    memo["sections"]["funds"]["rows"][1][2] = "n/a"
    funds = next(b for b in vip_memo_v2_export.build_blocks(memo) if b[0] == "table" and b[1][0] == "Category")
    assert funds[2][-1][2] == ""
    memo["sections"]["funds"]["rows"][1][2] = "30%"
    funds = next(b for b in vip_memo_v2_export.build_blocks(memo) if b[0] == "table" and b[1][0] == "Category")
    assert funds[2][-1][2] == "100%"


def test_export_special_glyphs_do_not_crash():
    memo = copy.deepcopy(FIXTURE)
    memo["sections"]["what"]["paragraphs"] = ["Roughly ≈ 5 → 10, ≤ 3 and ≥ 2 – ₹5 L, ₹ 7 Cr, 日本"]
    pdf_text, docx_text = _both_texts(memo)
    assert "~ 5 -> 10, <= 3 and >= 2 - Rs 5 L, Rs 7 Cr" in pdf_text.replace("\n", " ")
    assert "≈" in docx_text  # docx keeps real glyphs


def test_export_has_no_recommendation_text():
    verdict = FIXTURE["recommendation"]["verdict"]
    for text in _both_texts(FIXTURE):
        assert verdict not in text
        assert FIXTURE["recommendation"]["summary"] not in text
        assert FIXTURE["recommendation"]["conditions"][0] not in text


# ─── representation fixes (market cards, percents, empty questions, glyphs) ──

TBC = "[To be confirmed]"


def _shaped_memo() -> dict:
    memo = copy.deepcopy(FIXTURE)
    memo["sections"]["market"]["rows"] = [
        {"segment": "Cold rooms", "global_size": "$4.2B (2025, global)", "slice_low_cr": 230, "slice_high_cr": 780,
         "slice_label": "~₹230–780 Cr", "rationale": ["Storage held 68.0% (2025).", "Food share 84%."],
         "source": "Deck_founder@gmail.com.pdf; Fake Research 2026"},
        {"segment": "Pharma chain", "global_size": TBC, "slice_low_cr": 0, "slice_high_cr": 0,
         "slice_label": TBC, "rationale": "No third-party figure.", "source": TBC},
    ]
    memo["sections"]["market"]["total"] = [TBC, "Combined slice not stated."]
    memo["sections"]["funds"]["rows"] = [
        ["Deploy", "Sensors", "~14–17%", "₹70 L"],
        ["Models", ["ML engineers", "Compute"], "24%", "₹60 L"],
        ["Validation", "Lab", 15, "₹30 L"],
        ["Other", "Gap", TBC, TBC],
    ]
    memo["sections"]["team"]["holders"] = [["A", "Founder", 60, True], ["B", "Founder", TBC, False],
                                           ["C", "Investor", "CCD", TBC], ["D", "ESOP", None, False]]
    return memo


def test_export_market_is_per_segment_with_bullet_rationale():
    memo = _shaped_memo()
    blocks = vip_memo_v2_export.build_blocks(memo)
    i = blocks.index(("h1", "Addressable market"))
    j = next(n for n in range(i + 1, len(blocks)) if blocks[n][0] == "h1")
    market = blocks[i:j]
    assert ("h2", "Cold rooms") in market and ("h2", "Pharma chain") in market
    stats = [b for b in market if b[0] == "table"]
    assert stats[0][1] == ["Global market", "Realistic slice"]
    assert stats[0][2] == [["$4.2B (2025, global)", "~₹230–780 Cr"]]
    assert ("bullets", ["Storage held 68.0% (2025).", "Food share 84%."]) in market
    doc = Document(io.BytesIO(vip_memo_v2_export.render_docx_v2(memo)))
    bullets = [p.text for p in doc.paragraphs if p.style.name.startswith("List Bullet")]
    assert "Storage held 68.0% (2025)." in bullets and "Combined slice not stated." in bullets
    for text in _both_texts(memo):
        assert "globally" not in text
        assert ".," not in text
        assert "@gmail.com" not in text
        assert "Realistic opportunity, per year" in text


def test_export_percent_cells_keep_text_and_never_double_percent():
    memo = _shaped_memo()
    blocks = vip_memo_v2_export.build_blocks(memo)
    funds = next(b for b in blocks if b[0] == "table" and b[1][0] == "Category")
    assert [r[2] for r in funds[2]] == ["~14–17%", "24%", "15%", TBC, ""]
    holders = next(b for b in blocks if b[0] == "table" and b[1][0] == "Holder")
    assert [r[2] for r in holders[2]] == ["60%", TBC, "CCD", "—"]
    for text in _both_texts(memo):
        assert "%%" not in text and "]%" not in text and "—%" not in text
        assert "14-17%" in text or "14–17%" in text


def test_export_empty_questions_keeps_heading_with_neutral_line():
    memo = copy.deepcopy(FIXTURE)
    memo["questions"] = []
    for text in _both_texts(memo):
        assert "Questions for the founders" in text
        assert "No founder questions were generated." in text


def test_export_unknown_company_name_competitor_label():
    memo = copy.deepcopy(FIXTURE)
    memo["name"] = TBC
    blocks = vip_memo_v2_export.build_blocks(memo)
    comp = next(b for b in blocks if b[0] == "table" and b[1][0] == "Competitor")
    assert comp[1][3] == "Advantage"


def test_pdf_safe_glyphs_strip_diacritics_and_map_sub_superscripts():
    safe = vip_memo_v2_export._pdf_safe
    assert safe("CO₂ and O₂, TiO₂, x² ") == "CO2 and O2, TiO2, x2 "
    assert safe("−22 °C") == "-22 °C"
    assert safe("Krim-Nyāya, Pūjā, Skłodowska") == "Krim-Nyaya, Puja, Sklodowska"
    assert safe("café") == "café"  # cp1252 already has it
