"""Ingestion tests for the VIP memo generator (inputs + multi-format extraction)."""
from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from app.services.vip_memo_gen import inputs
from app.services.vip_memo_gen.extract import budget_texts, extract_any
from app.services.vip_memo_gen.inputs import FileRef


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Market Size"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "TAM is large"
    slide.notes_slide.notes_text_frame.text = "speaker secret note"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_text_and_notes():
    r = extract_any(_pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "d.pptx")
    assert r.skipped_reason is None
    assert "Slide 1:" in r.text
    assert "Market Size" in r.text and "TAM is large" in r.text
    assert "speaker secret note" in r.text


def test_xlsx_table_text():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Traction"
    ws.append(["Month", "Revenue"])
    ws.append(["Jan", 100])
    buf = io.BytesIO()
    wb.save(buf)
    r = extract_any(buf.getvalue(), "", "t.xlsx")
    assert "Traction" in r.text
    assert "Month\tRevenue" in r.text
    assert "Jan\t100" in r.text


def test_docx_tables_included():
    import docx

    d = docx.Document()
    d.add_paragraph("Intro paragraph")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Founder"
    t.cell(0, 1).text = "60%"
    buf = io.BytesIO()
    d.save(buf)
    r = extract_any(buf.getvalue(), "", "c.docx")
    assert "Intro paragraph" in r.text
    assert "Founder" in r.text and "60%" in r.text


def test_scanned_pdf_flags_transcription():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 700, "x")
    c.save()
    r = extract_any(buf.getvalue(), "application/pdf", "scan.pdf")
    assert r.needs_transcription is True


def test_text_pdf_not_flagged():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(30):
        c.drawString(40, 800 - i * 20, "A real line of searchable text content here. " * 3)
    c.save()
    r = extract_any(buf.getvalue(), "application/pdf", "t.pdf")
    assert r.needs_transcription is False
    assert "searchable" in r.text


def test_legacy_formats_skipped():
    for name in ("a.doc", "b.ppt", "c.xls"):
        r = extract_any(b"junk", "", name)
        assert r.skipped_reason == "unsupported_legacy_format"
    assert extract_any(b"junk", "", "z.zip").skipped_reason == "unsupported_type"


def test_images_need_transcription_and_csv_txt():
    r = extract_any(b"\x89PNG", "image/png", "a.png")
    assert r.needs_transcription and r.text == ""
    assert extract_any(b"a,b\n1,2", "text/csv", "a.csv").text == "a,b\n1,2"


def test_extract_never_raises_on_garbage():
    for name in ("a.pdf", "a.docx", "a.pptx", "a.xlsx"):
        r = extract_any(b"not a real file", "", name)
        assert r.skipped_reason == "extract_error"


def _ref(kind):
    return FileRef(bucket="b", path=f"{kind}.x", filename=f"{kind}.x", mime="text/plain", kind=kind)


def test_budget_priority_and_rollover():
    # cv listed first but lowest priority; deck is tiny so its share rolls over.
    items = [
        (_ref("cv"), "c" * 5000),
        (_ref("deck"), "d" * 10),
        (_ref("traction"), "t" * 5000),
    ]
    out = budget_texts(items, total_chars=1000)
    by = {r.kind: (t, trunc) for r, t, trunc in out}
    assert by["deck"] == ("d" * 10, False)
    assert by["traction"][1] is True and len(by["traction"][0]) > 10
    total = sum(len(t) for t, _ in by.values())
    assert total <= 1000
    assert [r.kind for r, _, _ in out][0] == "deck"
    # priority: traction gets at least as much as cv
    assert len(by["traction"][0]) >= len(by["cv"][0])


def test_application_block_has_no_pii():
    row = {
        "basic_email": "secret@x.com", "basic_phone": "+919999900000",
        "basic_full_name": "Hidden Person", "basic_name": "Hidden Person",
        "founder_email": "f@x.com", "contact_phone": "123",
        "basic_org": "Acme Robotics",
        "basic_teammates": [{"fullName": "Tia Team", "email": "tia@x.com", "role": "CTO", "phone": "55"}],
        "sip_founders": [{"name": "Fay Founder", "role": "CEO", "percent": 60, "email": "fay@x.com"}],
        "problem_describe": "Big problem",
        "sip_demo_video_url": "https://v.example/demo",
    }
    t = inputs.build_application_block(row, "sip")
    for bad in ("secret@x.com", "9999900000", "Hidden Person", "f@x.com", "tia@x.com", "fay@x.com", "@x.com"):
        assert bad not in t
    assert "Acme Robotics" in t and "Tia Team" in t and "CTO" in t
    assert "Fay Founder" in t and "60" in t
    assert "https://v.example/demo" in t


class _Q:
    def __init__(self, rows):
        self.rows = rows

    def __getattr__(self, _):
        return lambda *a, **k: self

    def execute(self):
        return SimpleNamespace(data=self.rows)


class _Client:
    def __init__(self, rows):
        self.rows = rows

    def table(self, _):
        return _Q(self.rows)


def test_collect_files_kinds(monkeypatch):
    from app.services import applications_query as aq

    monkeypatch.setattr(aq, "get_admin_client", lambda: _Client([]))
    row = {
        "sip_pitch_deck": {"path": "u/deck.pdf", "original_filename": "deck.pdf", "mime_type": "application/pdf"},
        "sip_cap_table_file": {"path": "u/cap.xlsx", "original_filename": "cap.xlsx"},
        "sip_traction_files": [{"path": "u/loi.pdf", "original_filename": "loi.pdf"}],
    }
    refs = inputs.collect_files(_Client([]), row, "sip")
    got = {r.kind: r for r in refs}
    assert set(got) == {"deck", "cap_table", "traction"}
    assert got["deck"].bucket == "sip-evidence-files" and got["deck"].mime == "application/pdf"
    assert got["cap_table"].mime.endswith("sheet")
    assert got["traction"].path == "u/loi.pdf"


def test_collect_files_resume_and_template(monkeypatch):
    from app.services import applications_query as aq

    monkeypatch.setattr(aq, "get_admin_client", lambda: _Client([{
        "storage_path": "u/cv.pdf", "original_filename": "cv.pdf",
        "file_size_bytes": 1, "mime_type": "application/pdf",
    }]))
    tpl = [{"storage_path": "u/t.docx", "original_filename": "t.docx", "mime_type": None}]
    row = {"resume_file_id": "r1", "user_id": "u"}
    refs = inputs.collect_files(_Client(tpl), row, "sip")
    kinds = {r.kind: r for r in refs}
    assert kinds["cv"].bucket == "sip-resumes"
    assert kinds["template"].bucket == "sip-application-templates"
    assert kinds["template"].mime.endswith("wordprocessingml.document")
    # no user id -> templates skipped
    assert "template" not in {r.kind for r in inputs.collect_files(_Client(tpl), {}, "sip")}


def test_budget_second_pass_gives_leftover_to_deck():
    items = [(_ref("deck"), "d" * 40000), (_ref("traction"), "t" * 1000), (_ref("cv"), "c" * 1000)]
    out = {r.kind: (t, trunc) for r, t, trunc in budget_texts(items, total_chars=60000)}
    assert len(out["deck"][0]) == 40000 and out["deck"][1] is False
    assert not out["traction"][1] and not out["cv"][1]


def test_pptx_title_once_and_grouped_text():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "UniqueTitle"
    grp = slide.shapes.add_group_shape()
    grp.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "InsideGroup"
    buf = io.BytesIO()
    prs.save(buf)
    t = extract_any(buf.getvalue(), "", "g.pptx").text
    assert t.count("UniqueTitle") == 1
    assert "InsideGroup" in t


def test_legacy_mimes_skipped_regardless_of_extension():
    for m in ("application/msword", "application/vnd.ms-powerpoint", "application/vnd.ms-excel"):
        assert extract_any(b"x", m, "file.bin").skipped_reason == "unsupported_legacy_format"


@pytest.mark.parametrize("phone", [
    "+91 98765 43210", "+91-9876543210", "+919876543210", "98765 43210", "9876543210",
    "09876543210", "(080) 2345 6789", "080-23456789", "+1 (415) 555-0123",
    "+44 20 7946 0958", "415-555-0123",
])
def test_scrub_contacts_removes_phones(phone):
    from app.services.vip_memo_gen.extract import scrub_contacts

    out = scrub_contacts(f"Call {phone} today.")
    assert out == "Call [phone removed] today."


def test_scrub_contacts_removes_emails():
    from app.services.vip_memo_gen.extract import scrub_contacts

    assert scrub_contacts("Mail ceo.name+vip@acme-robotics.co.in or a_b@x.io.") == (
        "Mail [email removed] or [email removed].")


@pytest.mark.parametrize("text", [
    "Raised ₹1,50,000 in 2024", "Ask: Rs 1,50,00,000 over 2019-2024", "TRL 6 to TRL-7",
    "FY2024-25 revenue ₹ 2.5 Cr", "10,000 units sold", "ISO 9001:2015", "1.5 million users",
    "Patent application 202141012345", "Pincode 560012", "Grant of $250,000 (2023)",
    "Revenue 12,34,56,789 INR",
])
def test_scrub_contacts_keeps_years_amounts_trl(text):
    from app.services.vip_memo_gen.extract import scrub_contacts

    assert scrub_contacts(text) == text
