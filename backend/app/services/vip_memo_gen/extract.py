"""Multi-format text extraction for VIP memo inputs. Never raises."""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass

log = logging.getLogger(__name__)

# Priority order (highest first) and relative weights for budget_texts.
_PRIORITY_WEIGHTS = [
    ("deck", 4), ("template", 3), ("traction", 3), ("milestones", 2),
    ("cap_table", 2), ("patents", 1), ("cv", 1),
]
PRIORITY = [k for k, _ in _PRIORITY_WEIGHTS]
_WEIGHTS = dict(_PRIORITY_WEIGHTS)
_LEGACY_MIMES = {"application/msword", "application/vnd.ms-powerpoint", "application/vnd.ms-excel"}

_SCANNED_PDF_CHARS_PER_PAGE = 200
_XLSX_MAX_ROWS = 200
_LEGACY = {"doc", "ppt", "xls"}


@dataclass
class Extracted:
    text: str
    method: str
    needs_transcription: bool = False
    skipped_reason: str | None = None


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = "\n\n".join(pages).strip()
    total = sum(len(p.strip()) for p in pages)
    scanned = not pages or total / len(pages) < _SCANNED_PDF_CHARS_PER_PAGE
    return Extracted(text, "pypdf", needs_transcription=scanned)


def _docx(data: bytes) -> Extracted:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return Extracted("\n".join(parts), "python-docx")


def _shape_lines(shapes, skip_id, lines: list[str]) -> None:
    for shape in shapes:
        if shape.shape_id == skip_id:
            continue
        if hasattr(shape, "shapes"):  # GroupShape
            _shape_lines(shape.shapes, skip_id, lines)
        elif shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                lines.append(t)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                lines.append("\t".join(c.text.strip() for c in row.cells))


def _pptx(data: bytes) -> Extracted:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None
        if title_shape is not None and title_shape.text_frame.text.strip():
            lines.append(title_shape.text_frame.text.strip())
        _shape_lines(slide.shapes, title_id, lines)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        out.append(f"Slide {i}: " + "\n".join(lines))
    return Extracted("\n\n".join(out), "python-pptx")


def _xlsx(data: bytes) -> Extracted:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"Sheet: {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= _XLSX_MAX_ROWS:
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    return Extracted("\n".join(parts), "openpyxl")


def extract_any(data: bytes, mime: str, filename: str) -> Extracted:
    ext = _ext(filename or "")
    mime = (mime or "").lower()
    try:
        if ext in _LEGACY or mime in _LEGACY_MIMES:
            return Extracted("", "none", skipped_reason="unsupported_legacy_format")
        if ext == "pdf" or mime == "application/pdf":
            return _pdf(data)
        if ext == "docx" or mime.endswith("wordprocessingml.document"):
            return _docx(data)
        if ext == "pptx" or mime.endswith("presentationml.presentation"):
            return _pptx(data)
        if ext == "xlsx" or mime.endswith("spreadsheetml.sheet"):
            return _xlsx(data)
        if ext in ("csv", "txt") or mime in ("text/csv", "text/plain"):
            return Extracted(data.decode("utf-8", errors="replace"), "text")
        if ext in ("png", "jpg", "jpeg") or mime in ("image/png", "image/jpeg"):
            return Extracted("", "image", needs_transcription=True)
        return Extracted("", "none", skipped_reason="unsupported_type")
    except Exception:
        log.warning("vip_memo_gen extract failed", extra={"file_ext": ext}, exc_info=True)
        return Extracted("", "none", skipped_reason="extract_error")


_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
# Separators inside a number are a single space, dot or dash, never a tab or newline,
# because DOCX/XLSX table cells are joined with tabs. Each alternative is anchored so it
# never starts or ends inside a longer number, keeping years, TRL levels, grouped amounts
# (1,50,000) and patent numbers intact. An unprefixed Indian mobile must be one contiguous
# 10-digit run: "90000 12000" is far more often two table figures than a phone number.
_PHONE = re.compile(
    r"(?<![\w+])\+\d{1,3}(?:[ .-]?\(\d{1,4}\))?(?:[ .-]?\d{2,5}){1,4}(?![\d,])"  # +91 98765 43210
    r"|(?<![\d,.])\(\d{2,5}\)[ .-]?\d{3,4}[ .-]?\d{3,4}(?![\d,])"               # (080) 2345 6789
    r"|(?<![\w,.])(?:91|0)[ -]?[6-9]\d{4}[ -]?\d{5}(?![\d,])"                     # 098450 12345
    r"|(?<![\d,.₹$])[6-9]\d{9}(?![\d,])"                                          # 9876543210
    r"(?!\s*(?:INR|USD|Rs\b|rupees|crores?|lakhs?|cr\b))"
    r"|(?<![\d,.])0\d{2,4}[ -]\d{3,4}[ -]?\d{3,4}(?![\d,])"                        # 080 2345 6789
    r"|(?<![\d,.])\d{3}[.-]\d{3}[.-]\d{4}(?![\d,])",                              # 415-555-0123
    re.IGNORECASE,
)


def _phone(m: re.Match) -> str:
    digits = sum(c.isdigit() for c in m.group(0))
    return "[phone removed]" if 8 <= digits <= 15 else m.group(0)


def scrub_contacts(text: str) -> str:
    """Remove e-mail addresses and phone numbers from founder document text (spec D6)."""
    return _PHONE.sub(_phone, _EMAIL.sub("[email removed]", text or ""))


def budget_texts(items, total_chars: int = 60000):
    """Allocate ``total_chars`` across files by priority; unused share rolls on.

    Returns ``[(ref, text, truncated)]`` in priority order.
    """
    def rank(kind: str) -> int:
        return PRIORITY.index(kind) if kind in PRIORITY else len(PRIORITY)

    ordered = sorted(items, key=lambda it: rank(it[0].kind))  # stable
    remaining = total_chars
    used: list[str] = []
    for idx, (ref, text) in enumerate(ordered):
        rest = sum(_WEIGHTS.get(r.kind, 1) for r, _ in ordered[idx:])
        share = remaining * _WEIGHTS.get(ref.kind, 1) // rest
        used.append(text[:share])
        remaining -= len(used[-1])
    # Second pass: leftover budget goes to truncated items, highest priority first.
    for i, (_, text) in enumerate(ordered):
        if remaining <= 0:
            break
        extra = min(remaining, len(text) - len(used[i]))
        if extra > 0:
            used[i] = text[: len(used[i]) + extra]
            remaining -= extra
    return [(ref, u, len(u) < len(text)) for (ref, text), u in zip(ordered, used, strict=False)]
